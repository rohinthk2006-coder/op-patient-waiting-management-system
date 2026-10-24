import os
import sys
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# --- Custom Numbered Canvas for PDF ---
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_header_footer(num_pages)
            super().showPage()
        super().save()

    def draw_header_footer(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Header (Pages > 1)
        if self._pageNumber > 1:
            self.drawString(40, letter[1] - 30, "Hospital OP Live Tracker — Market Requirements Document (MRD)")
            self.drawRightString(letter[0] - 40, letter[1] - 30, "Confidential • Healthcare Technology Systems")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(40, letter[1] - 34, letter[0] - 40, letter[1] - 34)
            
        # Footer
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(40, 42, letter[0] - 40, 42)
        
        self.drawString(40, 30, "Hospital OP Live Tracker v1.0.0 • Product Architecture & Market Requirements")
        self.drawRightString(letter[0] - 40, 30, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()


def set_cell_margins_and_bg(cell, fill_hex=None, top=100, bottom=100, left=140, right=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    if fill_hex:
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), fill_hex)
        tcPr.append(shd)
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def build_mrd_docx(filename):
    doc = docx.Document()
    
    # Page Setup: Letter, 0.6 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)
        s.different_first_page_header_footer = True
        
        # Header for page 2+
        header = s.header
        hp = header.paragraphs[0]
        hp.text = "Hospital OP Live Tracker — Market Requirements Document (MRD)"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(100, 116, 139)
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.text = "Hospital OP Live Tracker v1.0.0 • Baseline Approved • Document Deliverable A"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.size = Pt(8.5)
        fp.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(15, 23, 42)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # Document Header Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_kicker = p_title.add_run("MARKET REQUIREMENTS DOCUMENT (MRD) • DELIVERABLE A\n")
    r_kicker.font.size = Pt(9.5)
    r_kicker.font.bold = True
    r_kicker.font.color.rgb = RGBColor(13, 148, 136) # Teal
    
    r_title = p_title.add_run("Hospital OP Live Tracker")
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Multi-Hospital Outpatient Queue Orchestration, AI Waiting-Time Intelligence & Emergency Bed Tracking System")
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    
    # Metadata Table
    meta_table = doc.add_table(rows=3, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    col_widths = [Inches(1.3), Inches(2.3), Inches(1.3), Inches(2.4)]
    
    metadata_data = [
        [("Document Version:", True), ("1.0.0 (Approved Baseline)", False), ("Review Status:", True), ("Formal Baseline Approved", False)],
        [("Release Date:", True), ("October 24, 2026", False), ("Author / Owner:", True), ("Principal PM & Lead Healthcare Architect", False)],
        [("Geographic Pilot:", True), ("Tiruchirappalli (Trichy), TN, India", False), ("Target Stack:", True), ("React 19 + FastAPI + MongoDB + WebSockets", False)]
    ]
    
    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            label, is_bold = metadata_data[row_idx][col_idx]
            cell.text = label
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            if is_bold:
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(30, 58, 138)
                set_cell_margins_and_bg(cell, "F1F5F9", 80, 80, 100, 100)
            else:
                p.runs[0].font.color.rgb = RGBColor(51, 65, 85)
                set_cell_margins_and_bg(cell, "FFFFFF", 80, 80, 100, 100)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_section_heading(title, number=""):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(f"{number} {title}".strip())
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        
    def add_subheading(title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(title)
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(13, 148, 136)

    # 1. Executive Summary & Product Vision
    add_section_heading("Executive Summary & Product Vision", "1.")
    p = doc.add_paragraph()
    p.add_run("Outpatient (OP) departments across public and private hospitals face severe operational bottlenecks characterized by chaotic waiting rooms, opaque queue sequencing, predatory token hoarding, high patient anxiety, and critical specialist time wastage. The ").font.color.rgb = RGBColor(15, 23, 42)
    p.add_run("Hospital OP Live Tracker").bold = True
    p.add_run(" is an enterprise healthcare platform engineered to eliminate waiting room congestion through a single authoritative queue engine combining online token pre-booking with walk-in triage, automated arrival check-ins, autonomous background token expiry, high-contrast digital waiting hall signage, transparent waiting-time intelligence, and real-time emergency bed availability tracking.")

    # 2. Problem Statement & Current OP Pain Points
    add_section_heading("Problem Statement & Outpatient Operational Pain Points", "2.")
    pain_points = [
        ("Chaotic Waiting Halls & Nosocomial Risk: ", "Patients congregate in dense physical waiting areas for 2 to 4 hours without visibility into when they will be examined, elevating disease transmission risks and patient-staff conflicts."),
        ("Queue Fraud & Duplicate Bookings: ", "Unregulated manual paper queues or unindexed booking apps suffer from duplicate token issuance, queue-jumping, and token hoarding by scalpers."),
        ("Specialist Slot Wastage via Unmanaged No-Shows: ", "Patients who book tokens but abandon visits without notifying staff waste 15% to 30% of prime consulting slots because traditional systems lack autonomous expiry mechanisms to advance the queue."),
        ("Opaque Doctor Delays: ", "When emergency surgeries, rounds, or urgent trauma cases delay a doctor, waiting outpatients receive zero notice, magnifying frustration and crowd agitation."),
        ("Disconnected Emergency Bed Intelligence: ", "During regional medical emergencies, patients and 108 ambulance crews waste golden-hour transit time visiting hospitals blindly because real-time ICU and oxygen bed availability is not publicly visible.")
    ]
    for bold_prefix, text in pain_points:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.color.rgb = RGBColor(30, 58, 138)
        p.add_run(text)

    # 3. Target Market & Intended Hospital Segments
    add_section_heading("Target Market & Intended Hospital Segments", "3.")
    p = doc.add_paragraph()
    p.add_run("The product addresses healthcare delivery institutions across India, with an initial regional validation anchor in ").font.color.rgb = RGBColor(15, 23, 42)
    p.add_run("Tiruchirappalli (Trichy) District, Tamil Nadu").bold = True
    p.add_run(", spanning three distinct operational segments:")

    segments = [
        ("Segment 1: Government General & Teaching Medical Colleges (e.g., MGMGH Trichy, SRM Medical College): ", "Characterized by extreme patient volume (1,500–3,500 OP visits/day), predominantly walk-in registrations, requiring rugged zero-duplicate token generation, dual-tone audio chimes, and high-contrast digital signage."),
        ("Segment 2: Regional Multi-Specialty Private Hospital Networks (e.g., Kauvery Cantonment & Heartcity, Apollo Speciality, Frontline Hospital): ", "Driven by patient experience, slot punctuality, mobile OTP authentication, and transparent wait-time forecasting to reduce waiting lounge crowding."),
        ("Segment 3: Charitable & Mission Hospitals (e.g., CSI Mission Hospital): ", "Requiring low-friction technology adoption with dual-engine offline resilience and shared emergency bed reporting.")
    ]
    for bold_prefix, text in segments:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        p.add_run(text)

    # 4. Stakeholder Analysis & Adoption Barriers
    add_section_heading("Stakeholder Groups, Goals & Adoption Barriers", "4.")
    
    stk_table = doc.add_table(rows=7, cols=4)
    stk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    stk_widths = [Inches(1.2), Inches(2.0), Inches(2.0), Inches(2.1)]
    stk_data = [
        ["Stakeholder", "Primary Goal", "Key Functional Need", "Adoption Barrier & Mitigation"],
        ["Patient & Caregiver", "Minimize physical wait time; know exact queue status", "Passwordless OTP booking, live queue position, travel alerts", "Digital literacy: simplified UI, walk-in parity, audio chimes"],
        ["OP Registration Staff", "Process walk-in crowds without duplicate errors", "Fast walk-in desk booking, automated check-in confirmation", "Workload friction: single-click 5-second walk-in issuance"],
        ["Nurses & OPD Staff", "Maintain orderly queue flow outside consulting rooms", "Queue status screen, patient attendance tracking, hold queue", "Cluttered software: clean tablet-optimized queue manager"],
        ["Consulting Doctor", "Focus on patient care without administrative friction", "Single-click console (Call, Start, Complete), delay broadcast", "Time loss: 1-tap delay broadcast (+15m/+30m) with live push"],
        ["Hospital Admin", "Maximize room utilization; reduce waiting room chaos", "Live throughput analytics, bed census update, policy config", "System downtime: dual-engine offline memory store resilience"],
        ["Platform Admin", "Oversee regional health network and tenant integrity", "Multi-tenant onboarding, cross-hospital isolation, health API", "Security concerns: role-based access control, tenant scoping"]
    ]
    for row_idx, row in enumerate(stk_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = stk_widths[col_idx]
            cell.text = stk_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8)
            if row_idx == 0:
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                set_cell_margins_and_bg(cell, "1E3A8A", 80, 80, 100, 100)
            else:
                bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                set_cell_margins_and_bg(cell, bg, 60, 60, 80, 80)

    # 5. Value Propositions
    add_section_heading("Value Propositions by Stakeholder", "5.")
    val_props = [
        ("For Patients: ", "Zero wasted hospital hours. Book from home, receive travel departure recommendations based on traffic, view live queue movement, and avoid arriving hours before the doctor is ready."),
        ("For Doctors: ", "Continuous, orderly patient throughput. No unannounced walk-ins barging into consulting rooms. Single-click delay broadcasting automatically adjusts all patient reporting deadlines when emergency surgeries arise."),
        ("For Staff: ", "Elimination of queue arguments. Server-authoritative token numbers, automated auto-expiry for no-shows, and digital waiting hall signage ensure fair, transparent patient calling."),
        ("For Administrators: ", "Actionable throughput velocity metrics, recovered appointment slots, optimized staff allocation, and regional emergency bed leadership.")
    ]
    for bold_prefix, text in val_props:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.color.rgb = RGBColor(13, 148, 136)
        p.add_run(text)

    # 6. Market-Level Requirements Table
    add_section_heading("Market-Level Product Requirements & Priorities", "6.")
    
    mrd_table = doc.add_table(rows=15, cols=5)
    mrd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    mrd_widths = [Inches(0.8), Inches(2.2), Inches(2.3), Inches(0.9), Inches(1.1)]
    mrd_data = [
        ["Req ID", "Market / User Need", "Proposed Product Capability", "Priority", "Validation Status"],
        ["MR-001", "Frictionless patient identity without passwords", "Mobile OTP authentication with auto-provisioning", "Must Have", "Verified in Repo"],
        ["MR-002", "Discover hospital services and doctor schedules", "Multi-hospital directory, doctor OPD schedules", "Must Have", "Verified in Repo"],
        ["MR-003", "Pre-book outpatient tokens with arrival deadlines", "Atomic online booking with reporting deadline", "Must Have", "Verified in Repo"],
        ["MR-004", "Equal access for walk-in and online outpatients", "Unified queue engine with zero duplicate token guarantee", "Must Have", "Verified in Repo"],
        ["MR-005", "Recover unused consultation slots from no-shows", "1-tap arrival check-in and automated background expiry", "Must Have", "Verified in Repo"],
        ["MR-006", "Consulting room flow and delay communication", "Doctor console (Call, Visit, Delay broadcast push)", "Must Have", "Verified in Repo"],
        ["MR-007", "Manage hospital departments, staff, and policies", "Hospital administration and OP schedule configuration", "Should Have", "Partially Implemented"],
        ["MR-008", "Platform-level multi-tenant health network", "Platform admin onboarding with tenant data isolation", "Must Have", "Planned (Spec Defined)"],
        ["MR-009", "Real-time critical bed visibility for ambulances", "Emergency bed tracker (ICU, Trauma, O2) with status", "Must Have", "Verified in Repo"],
        ["MR-010", "Transparent waiting time expectation management", "Transparent heuristic wait forecast + ML regression hooks", "Must Have", "Verified in Repo"],
        ["MR-011", "Instant status alerts across multiple channels", "Real-time WebSockets hub + multi-event SMS/in-app alerts", "Must Have", "Verified in Repo"],
        ["MR-012", "Accessible waiting hall calling for elderly/ill", "Waiting Hall TV mode signage with dual-tone audio chime", "Must Have", "Verified in Repo"],
        ["MR-013", "Travel guidance to prevent missing OP deadlines", "Trichy locality travel time calculation with rush hour buffer", "Should Have", "Verified Heuristic"],
        ["MR-014", "Operational visibility into clinic throughput", "Hourly outpatient arrival/completion velocity charts", "Should Have", "Verified in Repo"]
    ]
    for row_idx, row in enumerate(mrd_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = mrd_widths[col_idx]
            cell.text = mrd_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(7.5)
            if row_idx == 0:
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                set_cell_margins_and_bg(cell, "1E3A8A", 60, 60, 80, 80)
            else:
                bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                set_cell_margins_and_bg(cell, bg, 50, 50, 70, 70)

    # 7. Product Scope, Exclusions & Future Opportunities
    add_section_heading("Product Scope, Exclusions & Strategic Roadmap", "7.")
    add_subheading("In-Scope Functional Boundaries (Phase 1 & 2):")
    p = doc.add_paragraph()
    p.add_run("Multi-hospital directory; patient mobile OTP authentication; atomic online OP token reservation; walk-in OP desk registration; unified FIFO queue ordering; 1-tap arrival check-in; autonomous background expiry worker with crash recovery; doctor consultation console; doctor OPD delay broadcast push; high-contrast waiting hall TV signage with Web Audio chimes; baseline heuristic wait-time prediction; emergency bed status dashboard; and regional travel time calculation.")
    
    add_subheading("Explicit Exclusions (Out of Scope):")
    exclusions = [
        ("No Clinical Diagnosis / Medical Triage: ", "The system does not provide medical diagnostic recommendations, prescribe drugs, or substitute for emergency 108 triage."),
        ("No Inpatient / Billing Integration: ", "Does not process hospital inpatient ward admissions, complex insurance claim reimbursements, or dispensary inventory."),
        ("No Biometric Hardware Mandate: ", "Aadhaar e-KYC biometric fingerprint/iris scanners are excluded; system operates with phone OTP and masked identity inputs in compliance with Indian privacy laws.")
    ]
    for bold_prefix, text in exclusions:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)

    # 8. Success Indicators & Measurement Methods
    add_section_heading("Success Indicators & Target Metrics", "8.")
    kpi_table = doc.add_table(rows=6, cols=3)
    kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    kpi_widths = [Inches(2.2), Inches(2.2), Inches(2.9)]
    kpi_data = [
        ["Key Performance Indicator (KPI)", "Target Operational Benchmark", "Measurement & Verification Method"],
        ["Waiting Room Peak Crowd Density", "> 40% reduction in physical waiting crowd", "Pre- and post-pilot waiting lounge headcounts"],
        ["No-Show Consultation Slot Recovery", "> 70% of abandoned slots recycled to queue", "Audit log tracking of EXPIRED tokens advancing next patient"],
        ["Wait Time Forecast Accuracy", "Within ±15% of actual consultation call time", "Comparison of predicted wait minutes vs actual call timestamps"],
        ["Zero Duplicate Token Incidents", "100% mathematical integrity (0 duplicates)", "Compound database unique constraint verification & stress tests"],
        ["Emergency Bed Update Latency", "< 30 minutes from physical bed status change", "Audit timestamps of bed update API calls against hospital census"]
    ]
    for row_idx, row in enumerate(kpi_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = kpi_widths[col_idx]
            cell.text = kpi_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8)
            if row_idx == 0:
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                set_cell_margins_and_bg(cell, "1E3A8A", 60, 60, 80, 80)
            else:
                bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                set_cell_margins_and_bg(cell, bg, 50, 50, 70, 70)

    # 9. Product Roadmap & Phased Evolution
    add_section_heading("High-Level Product Roadmap", "9.")
    roadmap_items = [
        ("Phase 1: Foundation & Core Queue Prototype (Current Status): ", "Scaffolded React 19 frontend and FastAPI backend; verified in-memory store; verified atomic booking, check-in, background auto-expiry worker, TV display, and 11 automated pytest test cases."),
        ("Phase 2: Pilot Hospital Department Hardening (Weeks 1–4): ", "Deploy in single OPD room (Cardiology at Apollo/Kauvery Trichy); validate doctor console UX; connect live MongoDB Atlas cluster; integrate transactional SMS gateway with DLT registered templates."),
        ("Phase 3: Multi-Hospital & Emergency Bed Expansion (Weeks 5–8): ", "Onboard all 6 Trichy partner hospitals; operationalize emergency bed tracking with ward nurse mobile interface; deploy public waiting hall screens."),
        ("Phase 4: AI Optimization & ABDM Integration (Weeks 9–16): ", "Train scikit-learn regression models on collected consultation timestamps; integrate Ayushman Bharat Health Account (ABHA) patient identity; support WhatsApp Business conversational booking.")
    ]
    for bold_prefix, text in roadmap_items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        p.add_run(text)

    # 10. Stakeholder Validation & Pilot Evaluation Plan
    add_section_heading("Stakeholder Validation & Operational Pilot Plan", "10.")
    p = doc.add_paragraph()
    p.add_run("To ensure clinical safety and operational adoption, the platform follows a structured 3-tier validation cycle:")
    steps = [
        ("Stage 1: Synthetic Simulation: ", "Validate concurrency with 30+ simultaneous booking threads and test background worker catch-up after simulated server crashes (Verified in `test_concurrency_and_expiry.py`)."),
        ("Stage 2: Single-Clinic Shadow Pilot: ", "Deploy the Doctor Console and Waiting Hall TV screen in a single outpatient clinic for 5 consecutive days without modifying hospital administrative records to assess doctor adherence and chime clarity."),
        ("Stage 3: Full OP Department Pilot: ", "Activate live patient booking with SMS dispatch and walk-in integration across General Medicine and Cardiology departments with weekly stakeholder retrospectives.")
    ]
    for bold_prefix, text in steps:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)

    doc.save(filename)
    print(f"MRD DOCX generated successfully at: {filename}")


def build_mrd_pdf(filename):
    doc = SimpleDocTemplate(
        filename,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=44,
        bottomMargin=48
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Palette
    c_primary = colors.HexColor("#1E3A8A")   # Navy
    c_teal = colors.HexColor("#0D9488")      # Teal
    c_dark = colors.HexColor("#0F172A")      # Slate 900
    c_slate = colors.HexColor("#475569")     # Slate 600
    c_border = colors.HexColor("#CBD5E1")    # Slate 300
    c_bg_light = colors.HexColor("#F8FAFC")  # Slate 50
    
    title_kicker_style = ParagraphStyle(
        'MRDTitleKicker',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=c_teal,
        spaceAfter=2
    )
    
    title_style = ParagraphStyle(
        'MRDTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=c_primary,
        spaceAfter=3
    )
    
    subtitle_style = ParagraphStyle(
        'MRDSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=c_slate,
        spaceAfter=8
    )
    
    h1_style = ParagraphStyle(
        'MRDH1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=c_primary,
        spaceBefore=8,
        spaceAfter=3,
        keepWithNext=True
    )
    
    h2_style = ParagraphStyle(
        'MRDH2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        leading=12.5,
        textColor=c_teal,
        spaceBefore=5,
        spaceAfter=2,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'MRDBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=c_dark,
        spaceAfter=4
    )
    
    bullet_style = ParagraphStyle(
        'MRDBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=c_dark,
        leftIndent=12,
        firstLineIndent=-8,
        spaceAfter=2
    )
    
    th_style = ParagraphStyle(
        'MRDTH',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white
    )
    
    td_style = ParagraphStyle(
        'MRDTD',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.2,
        leading=9.2,
        textColor=c_dark
    )
    
    td_bold_style = ParagraphStyle(
        'MRDTDBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.2,
        leading=9.2,
        textColor=c_primary
    )

    story = []
    
    # Title Block
    story.append(Paragraph("MARKET REQUIREMENTS DOCUMENT (MRD) • DELIVERABLE A", title_kicker_style))
    story.append(Paragraph("Hospital OP Live Tracker", title_style))
    story.append(Paragraph("Multi-Hospital Outpatient Queue Orchestration, AI Waiting-Time Intelligence & Emergency Bed Tracking System", subtitle_style))
    
    # Metadata Table
    meta_data = [
        [
            Paragraph("<b>Document Version:</b>", td_bold_style),
            Paragraph("1.0.0 (Approved Baseline)", td_style),
            Paragraph("<b>Review Status:</b>", td_bold_style),
            Paragraph("Formal Baseline Approved", td_style)
        ],
        [
            Paragraph("<b>Release Date:</b>", td_bold_style),
            Paragraph("October 24, 2026", td_style),
            Paragraph("<b>Document Owner:</b>", td_bold_style),
            Paragraph("Principal PM & Healthcare Solutions Architect", td_style)
        ],
        [
            Paragraph("<b>Geographic Pilot:</b>", td_bold_style),
            Paragraph("Tiruchirappalli (Trichy), TN, India", td_style),
            Paragraph("<b>Target Stack:</b>", td_bold_style),
            Paragraph("React 19 + FastAPI + MongoDB + WebSockets", td_style)
        ]
    ]
    t_meta = Table(meta_data, colWidths=[100, 170, 100, 170])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#F1F5F9")),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor("#F1F5F9")),
        ('GRID', (0, 0), (-1, -1), 0.5, c_border),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 6))

    # 1. Executive Summary & Vision
    story.append(Paragraph("1. Executive Summary & Product Vision", h1_style))
    story.append(Paragraph(
        "Outpatient (OP) departments across public and private hospitals face severe operational bottlenecks characterized by chaotic waiting rooms, opaque queue sequencing, predatory token scalping, high patient anxiety, and specialist time wastage. The <b>Hospital OP Live Tracker</b> is an enterprise healthcare platform engineered to eliminate waiting room congestion through a single authoritative queue engine combining online token pre-booking with walk-in triage, automated arrival check-ins, autonomous background token expiry, high-contrast digital waiting hall signage, transparent waiting-time intelligence, and real-time emergency bed availability tracking.",
        body_style
    ))

    # 2. Problem Statement & Pain Points
    story.append(Paragraph("2. Problem Statement & Operational Pain Points", h1_style))
    pain_points = [
        ("Chaotic Waiting Halls & Nosocomial Risk: ", "Patients congregate in dense physical waiting areas for 2 to 4 hours without visibility into when they will be examined, elevating disease transmission risks and patient-staff conflicts."),
        ("Queue Fraud & Duplicate Bookings: ", "Unregulated manual paper queues or unindexed booking apps suffer from duplicate token issuance, queue-jumping, and token hoarding by scalpers."),
        ("Specialist Slot Wastage via Unmanaged No-Shows: ", "Patients who book tokens but abandon visits without notifying staff waste 15% to 30% of prime consulting slots because traditional systems lack autonomous expiry mechanisms to advance the queue."),
        ("Opaque Doctor Delays: ", "When emergency surgeries, rounds, or urgent trauma cases delay a doctor, waiting outpatients receive zero notice, magnifying frustration and crowd agitation."),
        ("Disconnected Emergency Bed Intelligence: ", "During regional medical emergencies, patients and 108 ambulance crews waste golden-hour transit time visiting hospitals blindly because real-time ICU and oxygen bed availability is not publicly visible.")
    ]
    for pfx, desc in pain_points:
        story.append(Paragraph(f"• <b>{pfx}</b>{desc}", bullet_style))

    # 3. Target Market & Intended Hospital Segments
    story.append(Paragraph("3. Target Market & Intended Hospital Segments", h1_style))
    story.append(Paragraph(
        "The product addresses healthcare delivery institutions across India, with an initial regional validation anchor in <b>Tiruchirappalli (Trichy) District, Tamil Nadu</b>, spanning three distinct operational segments:",
        body_style
    ))
    segments = [
        ("Segment 1: Government General & Teaching Medical Colleges (e.g., MGMGH Trichy, SRM Medical College): ", "Extreme patient volume (1,500–3,500 OP visits/day), predominantly walk-in registrations, requiring rugged zero-duplicate token generation, dual-tone audio chimes, and high-contrast digital signage."),
        ("Segment 2: Regional Multi-Specialty Private Hospital Networks (e.g., Kauvery Cantonment & Heartcity, Apollo Speciality, Frontline Hospital): ", "Driven by patient experience, slot punctuality, mobile OTP authentication, and transparent wait-time forecasting to reduce waiting lounge crowding."),
        ("Segment 3: Charitable & Mission Hospitals (e.g., CSI Mission Hospital): ", "Low-friction technology adoption with dual-engine offline resilience and shared emergency bed reporting.")
    ]
    for pfx, desc in segments:
        story.append(Paragraph(f"• <b>{pfx}</b>{desc}", bullet_style))

    # 4. Stakeholder Analysis
    story.append(Paragraph("4. Stakeholder Groups, Goals & Adoption Barriers", h1_style))
    stk_rows = [
        [Paragraph("Stakeholder", th_style), Paragraph("Primary Goal", th_style), Paragraph("Key Functional Need", th_style), Paragraph("Adoption Barrier & Mitigation", th_style)],
        [Paragraph("Patient & Caregiver", td_bold_style), Paragraph("Minimize physical wait time; know exact status", td_style), Paragraph("Passwordless OTP booking, live queue position, travel alerts", td_style), Paragraph("Digital literacy: simplified UI, walk-in parity, audio chimes", td_style)],
        [Paragraph("OP Registration Staff", td_bold_style), Paragraph("Process walk-in crowds without duplicate errors", td_style), Paragraph("Fast walk-in desk booking, automated check-in confirmation", td_style), Paragraph("Workload friction: single-click 5-second walk-in issuance", td_style)],
        [Paragraph("Nurses & OPD Staff", td_bold_style), Paragraph("Maintain orderly queue flow outside consulting rooms", td_style), Paragraph("Queue status screen, patient attendance tracking, hold queue", td_style), Paragraph("Cluttered software: clean tablet-optimized queue manager", td_style)],
        [Paragraph("Consulting Doctor", td_bold_style), Paragraph("Focus on patient care without administrative friction", td_style), Paragraph("Single-click console (Call, Start, Complete), delay broadcast", td_style), Paragraph("Time loss: 1-tap delay broadcast (+15m/+30m) with live push", td_style)],
        [Paragraph("Hospital Admin", td_bold_style), Paragraph("Maximize room utilization; reduce waiting room chaos", td_style), Paragraph("Live throughput analytics, bed census update, policy config", td_style), Paragraph("System downtime: dual-engine offline memory store resilience", td_style)],
        [Paragraph("Platform Admin", td_bold_style), Paragraph("Oversee regional health network and tenant integrity", td_style), Paragraph("Multi-tenant onboarding, cross-hospital isolation, health API", td_style), Paragraph("Security concerns: role-based access control, tenant scoping", td_style)]
    ]
    t_stk = Table(stk_rows, colWidths=[90, 140, 150, 160])
    t_stk.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_primary),
        ('GRID', (0, 0), (-1, -1), 0.5, c_border),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, c_bg_light]),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_stk)
    story.append(Spacer(1, 6))

    # 5. Value Propositions
    story.append(Paragraph("5. Value Propositions by Stakeholder", h1_style))
    val_props = [
        ("For Patients: ", "Zero wasted hospital hours. Book from home, receive travel departure recommendations based on traffic, view live queue movement, and avoid arriving hours before the doctor is ready."),
        ("For Doctors: ", "Continuous, orderly patient throughput. No unannounced walk-ins barging into consulting rooms. Single-click delay broadcasting automatically adjusts all patient reporting deadlines when emergency surgeries arise."),
        ("For Staff: ", "Elimination of queue arguments. Server-authoritative token numbers, automated auto-expiry for no-shows, and digital waiting hall signage ensure fair, transparent patient calling."),
        ("For Administrators: ", "Actionable throughput velocity metrics, recovered appointment slots, optimized staff allocation, and regional emergency bed leadership.")
    ]
    for pfx, desc in val_props:
        story.append(Paragraph(f"• <b>{pfx}</b>{desc}", bullet_style))

    # 6. Market Requirements Table
    story.append(Paragraph("6. Market-Level Product Requirements & Priorities", h1_style))
    mrd_rows = [
        [Paragraph("Req ID", th_style), Paragraph("Market / User Need", th_style), Paragraph("Proposed Product Capability", th_style), Paragraph("Priority", th_style), Paragraph("Validation Status", th_style)],
        [Paragraph("MR-001", td_bold_style), Paragraph("Frictionless patient identity without passwords", td_style), Paragraph("Mobile OTP authentication with auto-provisioning", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-002", td_bold_style), Paragraph("Discover hospital services and doctor schedules", td_style), Paragraph("Multi-hospital directory, doctor OPD schedules", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-003", td_bold_style), Paragraph("Pre-book outpatient tokens with arrival deadlines", td_style), Paragraph("Atomic online booking with reporting deadline", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-004", td_bold_style), Paragraph("Equal access for walk-in and online outpatients", td_style), Paragraph("Unified queue engine with zero duplicate token guarantee", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-005", td_bold_style), Paragraph("Recover unused consultation slots from no-shows", td_style), Paragraph("1-tap arrival check-in and automated background expiry", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-006", td_bold_style), Paragraph("Consulting room flow and delay communication", td_style), Paragraph("Doctor console (Call, Visit, Delay broadcast push)", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-007", td_bold_style), Paragraph("Manage hospital departments, staff, and policies", td_style), Paragraph("Hospital administration and OP schedule configuration", td_style), Paragraph("Should Have", td_style), Paragraph("Partially Implemented", td_style)],
        [Paragraph("MR-008", td_bold_style), Paragraph("Platform-level multi-tenant health network", td_style), Paragraph("Platform admin onboarding with tenant data isolation", td_style), Paragraph("Must Have", td_style), Paragraph("Planned (Spec Defined)", td_style)],
        [Paragraph("MR-009", td_bold_style), Paragraph("Real-time critical bed visibility for ambulances", td_style), Paragraph("Emergency bed tracker (ICU, Trauma, O2) with status", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-010", td_bold_style), Paragraph("Transparent waiting time expectation management", td_style), Paragraph("Transparent heuristic wait forecast + ML regression hooks", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-011", td_bold_style), Paragraph("Instant status alerts across multiple channels", td_style), Paragraph("Real-time WebSockets hub + multi-event SMS/in-app alerts", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-012", td_bold_style), Paragraph("Accessible waiting hall calling for elderly/ill", td_style), Paragraph("Waiting Hall TV mode signage with dual-tone audio chime", td_style), Paragraph("Must Have", td_style), Paragraph("Verified in Repo", td_bold_style)],
        [Paragraph("MR-013", td_bold_style), Paragraph("Travel guidance to prevent missing OP deadlines", td_style), Paragraph("Trichy locality travel time calculation with rush hour buffer", td_style), Paragraph("Should Have", td_style), Paragraph("Verified Heuristic", td_style)],
        [Paragraph("MR-014", td_bold_style), Paragraph("Operational visibility into clinic throughput", td_style), Paragraph("Hourly outpatient arrival/completion velocity charts", td_style), Paragraph("Should Have", td_style), Paragraph("Verified in Repo", td_bold_style)]
    ]
    t_mrd = Table(mrd_rows, colWidths=[55, 150, 175, 65, 95])
    t_mrd.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_primary),
        ('GRID', (0, 0), (-1, -1), 0.5, c_border),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, c_bg_light]),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(t_mrd)
    story.append(Spacer(1, 6))

    # 7. Scope & Exclusions
    story.append(Paragraph("7. Scope, Exclusions & Strategic Roadmap", h1_style))
    story.append(Paragraph(
        "<b>In-Scope:</b> Multi-hospital directory; patient mobile OTP authentication; atomic online OP token reservation; walk-in OP desk registration; unified FIFO queue ordering; 1-tap arrival check-in; autonomous background expiry worker with crash recovery; doctor consultation console; doctor OPD delay broadcast push; high-contrast waiting hall TV signage with Web Audio chimes; baseline heuristic wait-time prediction; emergency bed status dashboard; and regional travel time calculation.",
        body_style
    ))
    exclusions = [
        ("No Clinical Diagnosis / Medical Triage: ", "The system does not provide medical diagnostic recommendations, prescribe drugs, or substitute for emergency 108 triage."),
        ("No Inpatient / Billing Integration: ", "Does not process hospital inpatient ward admissions, complex insurance claim reimbursements, or dispensary inventory."),
        ("No Biometric Hardware Mandate: ", "Aadhaar e-KYC biometric fingerprint/iris scanners are excluded; system operates with phone OTP and masked identity inputs in compliance with Indian privacy laws.")
    ]
    for pfx, desc in exclusions:
        story.append(Paragraph(f"• <b>{pfx}</b>{desc}", bullet_style))

    # 8. Success Indicators
    story.append(Paragraph("8. Success Indicators & Target Metrics", h1_style))
    kpi_rows = [
        [Paragraph("Key Performance Indicator (KPI)", th_style), Paragraph("Target Operational Benchmark", th_style), Paragraph("Measurement & Verification Method", th_style)],
        [Paragraph("Waiting Room Peak Crowd Density", td_bold_style), Paragraph("> 40% reduction in physical waiting crowd", td_style), Paragraph("Pre- and post-pilot waiting lounge headcounts", td_style)],
        [Paragraph("No-Show Slot Recovery Rate", td_bold_style), Paragraph("> 70% of abandoned slots recycled to queue", td_style), Paragraph("Audit log tracking of EXPIRED tokens advancing next patient", td_style)],
        [Paragraph("Wait Time Forecast Accuracy", td_bold_style), Paragraph("Within ±15% of actual consultation call time", td_style), Paragraph("Comparison of predicted wait minutes vs actual call timestamps", td_style)],
        [Paragraph("Zero Duplicate Token Incidents", td_bold_style), Paragraph("100% mathematical integrity (0 duplicates)", td_style), Paragraph("Compound database unique constraint verification & stress tests", td_style)],
        [Paragraph("Emergency Bed Update Latency", td_bold_style), Paragraph("< 30 minutes from physical bed status change", td_style), Paragraph("Audit timestamps of bed update API calls against hospital census", td_style)]
    ]
    t_kpi = Table(kpi_rows, colWidths=[160, 160, 220])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_primary),
        ('GRID', (0, 0), (-1, -1), 0.5, c_border),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, c_bg_light]),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 6))

    # 9. Roadmap
    story.append(Paragraph("9. High-Level Product Roadmap", h1_style))
    roadmap_items = [
        ("Phase 1: Foundation & Core Queue Prototype (Current Status): ", "Scaffolded React 19 frontend and FastAPI backend; verified in-memory store; verified atomic booking, check-in, background auto-expiry worker, TV display, and 11 automated pytest test cases."),
        ("Phase 2: Pilot Hospital Department Hardening (Weeks 1–4): ", "Deploy in single OPD room (Cardiology at Apollo/Kauvery Trichy); validate doctor console UX; connect live MongoDB Atlas cluster; integrate transactional SMS gateway with DLT registered templates."),
        ("Phase 3: Multi-Hospital & Emergency Bed Expansion (Weeks 5–8): ", "Onboard all 6 Trichy partner hospitals; operationalize emergency bed tracking with ward nurse mobile interface; deploy public waiting hall screens."),
        ("Phase 4: AI Optimization & ABDM Integration (Weeks 9–16): ", "Train scikit-learn regression models on collected consultation timestamps; integrate Ayushman Bharat Health Account (ABHA) patient identity; support WhatsApp Business conversational booking.")
    ]
    for pfx, desc in roadmap_items:
        story.append(Paragraph(f"• <b>{pfx}</b>{desc}", bullet_style))

    # 10. Validation Plan
    story.append(Paragraph("10. Stakeholder Validation & Operational Pilot Plan", h1_style))
    steps = [
        ("Stage 1: Synthetic Simulation: ", "Validate concurrency with 30+ simultaneous booking threads and test background worker catch-up after simulated server crashes (Verified in test suite)."),
        ("Stage 2: Single-Clinic Shadow Pilot: ", "Deploy the Doctor Console and Waiting Hall TV screen in a single outpatient clinic for 5 consecutive days without modifying hospital administrative records to assess doctor adherence and chime clarity."),
        ("Stage 3: Full OP Department Pilot: ", "Activate live patient booking with SMS dispatch and walk-in integration across General Medicine and Cardiology departments with weekly stakeholder retrospectives.")
    ]
    for pfx, desc in steps:
        story.append(Paragraph(f"• <b>{pfx}</b>{desc}", bullet_style))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"MRD PDF generated successfully at: {filename}")


if __name__ == '__main__':
    out_dir = r"D:\op patient waiting management system\docs\product-docs"
    docx_path = os.path.join(out_dir, "Hospital_OP_Live_Tracker_MRD.docx")
    pdf_path = os.path.join(out_dir, "Hospital_OP_Live_Tracker_MRD.pdf")
    build_mrd_docx(docx_path)
    build_mrd_pdf(pdf_path)
